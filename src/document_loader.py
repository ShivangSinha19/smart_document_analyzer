from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class LoadedDocument:
    filename: str
    text: str


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def load_uploaded_file(uploaded_file) -> LoadedDocument:
    """Extract readable text from a Streamlit UploadedFile."""
    filename = uploaded_file.name
    extension = Path(filename).suffix.lower()
    content = uploaded_file.getvalue()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension}'. Supported types: {supported}")

    if extension == ".pdf":
        text = _extract_pdf_text(content)
    elif extension == ".docx":
        text = _extract_docx_text(content)
    else:
        text = _extract_plain_text(content)

    cleaned = _normalize_text(text)
    if not cleaned:
        raise ValueError(f"No readable text found in {filename}")

    return LoadedDocument(filename=filename, text=cleaned)


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                table_cells.append(cell.text)
    return "\n".join(paragraphs + table_cells)


def _extract_plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


def _normalize_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()

